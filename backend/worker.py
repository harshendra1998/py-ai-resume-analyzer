import json
import logging
import os
import re
import subprocess

import docx
import fitz  # PyMuPDF
import httpx
import pytesseract
from celery import Celery
from PIL import Image
from sqlalchemy import select, update

from db import ANALYSIS_KEY, cache, cvs, engine

celery = Celery("cv", broker=os.environ["REDIS_URL"])
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://ollama:11434")
LLM_MODEL = os.environ.get("LLM_MODEL") or "qwen2.5:3b"
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")  # optional fallback when Ollama is down or returns bad JSON
GROQ_MODEL = os.environ.get("GROQ_MODEL") or "openai/gpt-oss-20b"

PROMPT = """Extract candidate data from the resume below. Reply with JSON only, using exactly these keys:
{"name": string, "mobile": string, "email": string, "location": string (city, country),
 "skills": [string] (key technical and professional skills, at most 25),
 "experience_years": number (total professional experience in years),
 "current_title": string, "education": [string], "companies": [string],
 "summary": string (two sentences)}
Use null for anything not present in the resume. Do not invent data.

RESUME:
"""

JD_PROMPT = """Extract the hiring requirements from the job description below. Reply with JSON only, using exactly these keys:
{"title": string, "required_skills": [string] (must-have skills, at most 15, short names like "Python"),
 "nice_to_have_skills": [string] (optional or preferred skills, at most 10),
 "min_experience_years": number (minimum years of experience asked for, null if not stated)}

JOB DESCRIPTION:
"""


def ocr_page(page):
    pix = page.get_pixmap(dpi=300)
    return pytesseract.image_to_string(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))


def extract_text(path):
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        with fitz.open(path) as doc:
            text = "\n".join(p.get_text() for p in doc)
            if len(text.strip()) < 50:  # no text layer -> scanned PDF, OCR it
                text = "\n".join(ocr_page(p) for p in doc)
        return text
    if ext == "docx":
        d = docx.Document(path)
        parts = [p.text for s in d.sections for p in s.header.paragraphs]  # contact info often lives in the header
        parts += [p.text for p in d.paragraphs]
        parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
        return "\n".join(parts)
    if ext == "doc":
        return subprocess.run(["antiword", path], capture_output=True, text=True, check=True).stdout
    raise ValueError(f"unsupported file type: .{ext}")


def llm_json(prompt, timeout=600):
    """Ask the local Ollama model for JSON; fall back to Groq if it fails and a key is set."""
    messages = [{"role": "user", "content": prompt}]
    try:
        r = httpx.post(f"{OLLAMA_URL}/api/chat", timeout=timeout, json={
            "model": LLM_MODEL, "format": "json", "stream": False,
            "options": {"temperature": 0, "num_ctx": 8192}, "messages": messages})
        r.raise_for_status()
        return json.loads(r.json()["message"]["content"])
    except (httpx.HTTPError, ValueError, KeyError) as e:
        if not GROQ_API_KEY:
            raise
        logging.warning("Ollama failed (%s), falling back to Groq", e)
    r = httpx.post("https://api.groq.com/openai/v1/chat/completions", timeout=120,
                   headers={"Authorization": f"Bearer {GROQ_API_KEY}"},
                   json={"model": GROQ_MODEL, "temperature": 0, "response_format": {"type": "json_object"},
                         "messages": messages})
    r.raise_for_status()
    return json.loads(r.json()["choices"][0]["message"]["content"])


def llm_extract(text):
    return llm_json(PROMPT + text[:12000])  # ponytail: truncates very long CVs


def clean_list(raw, limit):
    """LLM list (or comma string) -> trimmed, case-insensitively de-duplicated strings."""
    if isinstance(raw, str):
        raw = raw.split(",")
    seen, out = set(), []
    for x in raw or []:
        x = str(x).strip()
        if x and x.lower() not in seen:
            seen.add(x.lower())
            out.append(x)
    return out[:limit]


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def to_row(d):
    """Normalise whatever the LLM returned into DB columns."""
    def s(k):
        v = d.get(k)
        return (str(v).strip() or None) if v is not None else None

    return {
        "name": s("name"), "mobile": s("mobile"), "email": s("email"), "location": s("location"),
        "skills": clean_list(d.get("skills"), 25), "experience_years": to_float(d.get("experience_years")),
        "details": {k: d.get(k) for k in ("current_title", "education", "companies", "summary")},
    }


# ---------- JD matching ----------

def jd_requirements(jd):
    d = llm_json(JD_PROMPT + jd[:8000], timeout=180)
    required = clean_list(d.get("required_skills"), 15)
    taken = {s.lower() for s in required}
    return {
        "title": d.get("title"),
        "required_skills": required,
        "nice_to_have_skills": [s for s in clean_list(d.get("nice_to_have_skills"), 10) if s.lower() not in taken],
        "min_experience_years": to_float(d.get("min_experience_years")),
    }


def norm(skill):
    """'React.js' -> 'react', 'Node JS' -> 'node', 'C++' -> 'c++'."""
    s = re.sub(r"[^a-z0-9+#]", "", skill.lower())
    return s[:-2] if s.endswith("js") and len(s) > 4 else s


def has_skill(skill, cv_skills, text):
    """Matches the CV's extracted skills, or the CV text itself (catches skills the LLM left out).
    Text search is skipped for 1-2 letter skills like 'Go' or 'R', which would match ordinary words."""
    if norm(skill) in cv_skills:
        return True
    return len(skill) >= 3 and re.search(rf"(?<![a-z0-9]){re.escape(skill.lower())}(?![a-z0-9])", text) is not None


def score_cv(req, cv):
    """Fit score 0-100: 60% must-have skills, 10% nice-to-have skills, 30% experience vs. required."""
    cv_skills, text = {norm(s) for s in cv["skills"]}, (cv["raw_text"] or "").lower()
    required, nice = req["required_skills"], req["nice_to_have_skills"]
    matched = [s for s in required if has_skill(s, cv_skills, text)]
    nice_matched = [s for s in nice if has_skill(s, cv_skills, text)]
    need, exp = req["min_experience_years"], cv["experience_years"]
    skill_fit = len(matched) / len(required) if required else 1.0
    nice_fit = len(nice_matched) / len(nice) if nice else 0.0
    exp_fit = 1.0 if not need else min((exp or 0) / need, 1.0)
    return {
        "id": cv["id"], "name": cv["name"], "email": cv["email"], "mobile": cv["mobile"],
        "location": cv["location"], "experience_years": exp, "filename": cv["filename"],
        "current_title": (cv["details"] or {}).get("current_title"),
        "score": round(100 * (0.6 * skill_fit + 0.1 * nice_fit + 0.3 * exp_fit)),
        "matched_skills": matched, "missing_skills": [s for s in required if s not in matched],
        "nice_matched": nice_matched,
    }


def top_matches(req, cvs_rows, n=5):
    """Best n distinct candidates that have at least one must-have skill (same person uploaded twice counts once)."""
    scored = sorted((score_cv(req, cv) for cv in cvs_rows),
                    key=lambda m: (m["score"], m["experience_years"] or 0), reverse=True)
    out, seen = [], set()
    for m in scored:
        who = (m["email"] or m["name"] or str(m["id"])).lower()
        if who in seen or (req["required_skills"] and not m["matched_skills"]):
            continue
        seen.add(who)
        out.append(m)
    return out[:n]


@celery.task
def process_cv(cv_id):
    with engine.begin() as c:
        path = c.execute(select(cvs.c.path).where(cvs.c.id == cv_id)).scalar_one()
        c.execute(update(cvs).where(cvs.c.id == cv_id).values(status="processing"))
    try:
        text = extract_text(path).replace("\x00", "")
        if not text.strip():
            raise ValueError("no text found in document")
        values = {**to_row(llm_extract(text)), "raw_text": text, "status": "done", "error": None}
    except Exception as e:
        values = {"status": "failed", "error": str(e)[:500]}
    with engine.begin() as c:
        c.execute(update(cvs).where(cvs.c.id == cv_id).values(**values))
    cache.delete(ANALYSIS_KEY)
